import hashlib
import re
from io import BytesIO

from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from sqlalchemy.orm import Session

from ..models import Document, DocumentChunk

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
INJECTION_PATTERNS = (
    r"ignore (all |any )?(previous|prior) instructions",
    r"system message",
    r"act as (an? )?administrator",
    r"reveal (the )?(prompt|secret|api key)",
    r"override (these|all) rules",
)


def parse_document(filename: str, content: bytes) -> list[tuple[str, str]]:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(415, "Only PDF and DOCX files are supported.")
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            return [(f"Page {index}", page.extract_text() or "") for index, page in enumerate(reader.pages, 1)]
        document = DocxDocument(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return [(f"Paragraph {index}", text) for index, text in enumerate(paragraphs, 1)]
    except Exception as exc:
        raise HTTPException(422, f"The document could not be parsed: {exc}") from exc


def detect_injection(text: str) -> list[str]:
    normalized = " ".join(text.lower().split())
    return [pattern for pattern in INJECTION_PATTERNS if re.search(pattern, normalized)]


def chunk_text(source_location: str, text: str, size: int = 1200) -> list[tuple[str, str]]:
    words = text.split()
    if not words:
        return []
    return [(source_location, " ".join(words[start : start + size])) for start in range(0, len(words), size)]


async def ingest_document(
    db: Session,
    upload: UploadFile,
    document_id: str,
    title: str,
    category: str,
    department: str | None,
    version: str,
    effective_date,
    expiry_date,
    max_upload_bytes: int,
) -> Document:
    content = await upload.read()
    if not content:
        raise HTTPException(422, "The uploaded document is empty.")
    if len(content) > max_upload_bytes:
        raise HTTPException(413, "The uploaded document exceeds the file-size limit.")
    filename = upload.filename or "upload"
    digest = hashlib.sha256(content).hexdigest()
    if db.query(Document).filter(Document.content_hash == digest).first():
        raise HTTPException(409, "Duplicate document content detected.")
    if db.query(Document).filter(Document.document_id == document_id, Document.version == version).first():
        raise HTTPException(409, "This document ID and version already exist.")

    parsed_pages = parse_document(filename, content)
    full_text = "\n".join(text for _, text in parsed_pages).strip()
    if not full_text:
        raise HTTPException(422, "No usable text was found in the document.")
    flags = detect_injection(full_text)
    db.query(Document).filter(Document.document_id == document_id).update({Document.is_active: False})
    document = Document(
        document_id=document_id,
        title=title,
        category=category,
        department=department,
        version=version,
        effective_date=effective_date,
        expiry_date=expiry_date,
        filename=filename,
        content_hash=digest,
        injection_flags=flags,
        is_active=not bool(flags),
        is_quarantined=bool(flags),
    )
    db.add(document)
    db.flush()
    chunk_number = 0
    for location, text in parsed_pages:
        for chunk_location, chunk in chunk_text(location, text):
            chunk_number += 1
            db.add(DocumentChunk(
                document_db_id=document.id,
                document_id=document_id,
                document_version=version,
                chunk_id=f"{document_id}-C{chunk_number:03d}",
                section_id=chunk_location.replace(" ", "-"),
                heading=None,
                source_location=chunk_location,
                text=chunk,
            ))
    db.commit()
    db.refresh(document)
    return document
